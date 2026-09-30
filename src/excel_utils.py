"""Helpers to write clean, analyst-ready Excel reports with openpyxl."""
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)


def write_sheet(writer, df, sheet_name, number_formats=None, color_scale_cols=None, reverse_scale=False):
    """Write a DataFrame to a formatted sheet.

    number_formats: {column_name: excel_format}, e.g. {"amount": "#,##0"}
    color_scale_cols: columns to shade green (good) -> red (bad)
    """
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]
    number_formats = number_formats or {}

    for col_idx, col_name in enumerate(df.columns, start=1):
        letter = get_column_letter(col_idx)
        header = ws.cell(row=1, column=col_idx)
        header.fill, header.font = HEADER_FILL, HEADER_FONT
        header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        fmt = number_formats.get(col_name)
        for row in range(2, len(df) + 2):
            cell = ws.cell(row=row, column=col_idx)
            cell.font = BODY_FONT
            if fmt:
                cell.number_format = fmt

        sample = [str(v) for v in df[col_name].head(200)]
        width = max([len(str(col_name))] + [len(s) for s in sample]) + 2
        ws.column_dimensions[letter].width = min(max(width, 10), 45)

        if color_scale_cols and col_name in color_scale_cols and len(df) > 1:
            low, high = ("F8696B", "63BE7B") if reverse_scale else ("63BE7B", "F8696B")
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{len(df) + 1}",
                ColorScaleRule(start_type="min", start_color=low,
                               mid_type="percentile", mid_value=50, mid_color="FFEB84",
                               end_type="max", end_color=high))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30
    return ws
